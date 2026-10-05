"""Gera a planilha de auditoria de termos (somente leitura: NÃO altera nenhum script).

Uso:  python auditoria_termos/gerar_auditoria.py
Saída: auditoria_termos/Auditoria_Termos.xlsx

Pareia cada linha de scripts_PT-BR com a linha de scripts_JP pelo mesmo label/coluna
e usa o termo JP como âncora pra classificar cada ocorrência:
  AUTO      -> JP confirma a troca
  REVISAR   -> ambíguo / JP não confirma / precisa de olho humano
  BLOQUEADO -> JP contradiz a troca (não trocar)
"""
import csv, hashlib, os, re, sys, urllib.parse
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

sys.stdout.reconfigure(encoding='utf-8')

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.path.join(REPO, 'Pasta principal')
PT_DIR = os.path.join(BASE, 'scripts_PT-BR')
JP_DIR = os.path.join(BASE, 'scripts_JP')
PROGRESSO = os.path.join(REPO, 'FATE REMASTERED PROGRESSO - Geral.csv')
OUT = os.path.join(REPO, 'auditoria_termos', 'Auditoria_Termos.xlsx')

AUTO, REV, BLOQ, DEC = 'AUTO', 'REVISAR', 'BLOQUEADO', 'DECISÃO'
TAG = re.compile(r'\[[^\]]*\]')


# ---------------------------------------------------------------- leitura DAT

ROW = re.compile(r'^\d+::')


def parse_dat(path):
    """{(label, coluna): (nº linha, texto)} para todas as colunas lstr do arquivo DAT,
    mais a lista de anomalias de formato [(nº linha, descrição, trecho)]."""
    out, anomalias = {}, []
    if not os.path.exists(path):
        return out, anomalias
    with open(path, encoding='utf-8', errors='replace', newline='') as fh:
        lines = fh.read().split('\n')
    if len(lines) < 2 or lines[0].strip('﻿\r') != 'DAT':
        return out, anomalias
    cols = [c.split('=') for c in lines[1].rstrip('\r').split('::') if c]
    last_row = max((i for i, ln in enumerate(lines, start=1) if ROW.match(ln)), default=0)
    for i, ln in enumerate(lines[2:], start=3):
        ln = ln.rstrip('\r')
        if not ln.strip() or i > last_row:  # depois da última linha só há lixo da decodificação
            continue
        if not ROW.match(ln):
            anomalias.append((i, 'linha solta, fora do formato id::label::texto::', ln[:200]))
            continue
        parts = ln.split('::')
        while parts and not parts[-1].strip():
            parts = parts[:-1]
        if len(parts) > len(cols):
            anomalias.append((i, f'{len(parts)} colunas, o cabeçalho tem {len(cols)} ("::" a mais no texto?)', ln[:200]))
        label = parts[1].strip('$') if len(parts) > 1 else parts[0]
        for ci, col in enumerate(cols):
            if len(col) == 2 and col[1] == 'lstr' and ci < len(parts):
                key = (label, col[0])
                if key in out:  # label duplicado: desambigua pelo id
                    key = (label + '#' + parts[0], col[0])
                out[key] = (i, parts[ci])
    return out, anomalias


def load_progresso():
    info = {}
    if not os.path.exists(PROGRESSO):
        return info
    with open(PROGRESSO, encoding='utf-8', errors='replace') as fh:
        for r in list(csv.reader(fh))[2:]:
            if len(r) > 4 and r[1].endswith('.epk_dec') and r[1] not in info:
                info[r[1]] = {'cena': r[0], 'revisores': r[2], 'estado': r[3], 'rota': r[4]}
    return info


# ---------------------------------------------------------------- utilidades de texto

CONNECT = {'de', 'do', 'da', 'dos', 'das', 'e', 'à', 'a'}


def title_pt(s):
    return ' '.join(w if w.lower() in CONNECT and k else w[:1].upper() + w[1:].lower()
                    for k, w in enumerate(s.split(' ')))


def keep_case(orig, new):
    """Mantém a inicial maiúscula/minúscula do original."""
    if orig[:1].isupper():
        return new[:1].upper() + new[1:]
    return new[:1].lower() + new[1:]


def in_tag(text, pos):
    op = text.rfind('[', 0, pos)
    return op != -1 and text.rfind(']', op, pos) == -1 and text.find(']', pos) != -1


def plain(s):
    return TAG.sub('', s)


def sentence_pos(text, pos, prev):
    """'inicio', 'meio' ou 'ambiguo' (reticências antes)."""
    s = plain(text[:pos]).rstrip(' 　\t"“”\'‘’「『(—–-*')
    if not s:
        if prev is None:
            return 'inicio'
        p = plain(prev).rstrip(' 　\t')
        p = p.rstrip('"“”\'’」』)*')
        if not p or p[-1] in '.!?…':
            return 'inicio'
        return 'meio'
    if s.endswith('...') or s.endswith('…'):
        return 'ambiguo'
    if s[-1] in '.!?':
        return 'inicio'
    return 'meio'


def jp_has(jp, rx):
    return bool(rx and jp is not None and re.search(rx, jp))


def classify(jp, ok, block, ok_desc, block_desc):
    if jp is None:
        return REV, 'linha sem par no JP'
    o, b = jp_has(jp, ok), jp_has(jp, block)
    if o and not b:
        return AUTO, f'JP confirma ({ok_desc})'
    if o and b:
        return REV, f'JP tem {ok_desc} e {block_desc} na mesma linha'
    if b:
        return BLOQUEADO_MSG(block_desc)
    return REV, f'JP não tem {ok_desc} nesta linha'


def BLOQUEADO_MSG(block_desc):
    return BLOQ, f'JP tem {block_desc} (não é o termo da regra)'


# ---------------------------------------------------------------- regras

class Rule:
    def __init__(self, rid, termo, obs, estado, pattern, flags=0, propose=None,
                 ok=None, block=None, ok_desc='', block_desc='', check=None, regra=''):
        self.rid, self.termo, self.obs, self.estado = rid, termo, obs, estado
        self.rx = re.compile(pattern, flags)
        self.propose = propose           # (m, text) -> (s, e, novo) | None p/ ignorar
        self.ok, self.block = ok, block
        self.ok_desc, self.block_desc = ok_desc or (ok or ''), block_desc or (block or '')
        self.check = check               # (m, text, ctx) -> (classe_forçada|None, motivo|None) | 'skip'
        self.regra = regra


def sub_span(repl_fn):
    """Proposta simples: troca só o trecho casado."""
    def f(m, text):
        new = repl_fn(m)
        if new is None or new == m.group(0):
            return None
        return m.start(), m.end(), new
    return f


def canon(fn):
    """Troca pra forma canônica; ignora quando já está canônico ou em CAIXA ALTA."""
    def g(m):
        t = m.group(0)
        if t.isupper():
            return None
        c = fn(t)
        return None if c == t else c
    return sub_span(g)


def plural(t, sing, plur):
    return plur if t.lower().rstrip().endswith('s') else sing


RULES = []
R = RULES.append

# --- Nomes: ordem japonesa (sobrenome + nome)
NOMES = [
    ('Rin', 'Tohsaka', r'遠坂\s*凛'), ('Sakura', 'Matou', r'間桐\s*桜'), ('Sakura', 'Tohsaka', r'遠坂\s*桜'),
    ('Shirou|Shiro', 'Emiya', r'衛宮\s*士郎'), ('Shinji', 'Matou', r'間桐\s*慎二'),
    ('Kiritsugu', 'Emiya', r'衛宮\s*切嗣'), ('Zouken|Zoken', 'Matou|Makiri', r'(間桐|マキリ)\s*臓硯'),
    ('Taiga', 'Fujimura', r'藤村\s*大河'), ('Kirei', 'Kotomine', r'言峰\s*綺礼'),
    ('Souichirou|Soichiro|Souichiro|Soichirou', 'Kuzuki', r'葛木\s*宗一郎'),
    ('Issei', 'Ryuudou|Ryudou|Ryuudo|Ryudo', r'柳洞\s*一成'), ('Ayako', 'Mitsuzuri', r'美綴\s*綾子'),
    ('Kane', 'Himuro', r'氷室\s*鐘'), ('Kaede', 'Makidera', r'蒔寺\s*楓'), ('Yukika', 'Saegusa', r'三枝\s*由紀香'),
    ('Tokiomi', 'Tohsaka', r'遠坂\s*時臣'), ('Aoi', 'Tohsaka', r'遠坂\s*葵'), ('Raiga', 'Fujimura', r'藤村\s*雷画'),
    ('Kariya', 'Matou', r'間桐\s*雁夜'), ('Byakuya', 'Matou', r'間桐\s*鶴野'),
]
for given, family, jp in NOMES:
    R(Rule('N01', 'Nomes (ordem JP)', '"Tohsaka Rin", e não "Rin Tohsaka" (vale pra todos os nomes)', 'Não Substituido',
           rf'\b({given}) ({family})\b', 0,
           propose=lambda m, t: (m.start(), m.end(), f'{m.group(2)} {m.group(1)}'),
           ok=jp, ok_desc=f'nome completo {jp.replace(chr(92) + "s*", "")}',
           regra='Nome ocidental → sobrenome + nome'))

# --- Sobrenomes sem plural
FAMILIAS = {'Tohsaka': '遠坂', 'Matou': '間桐', 'Emiya': '衛宮', 'Makiri': 'マキリ', 'Einzbern': 'アインツベルン',
            'Fujimura': '藤村', 'Kotomine': '言峰', 'Ryuudou': '柳洞', 'Mitsuzuri': '美綴', 'Himuro': '氷室',
            'Makidera': '蒔寺', 'Saegusa': '三枝', 'Kuzuki': '葛木'}
for fam, jp in FAMILIAS.items():
    R(Rule('N02', 'Makiri / sobrenomes no plural', 'Makiris viram Makiri. Também mexer em "Tohsakas" e "Matous", etc.',
           'Não Substituido', rf'\b({fam})s\b', 0,
           propose=lambda m, t: (m.start(), m.end(), m.group(1)), ok=jp,
           regra='Sobrenome no plural → singular'))


def compound_check(m, text, ctx):
    a, b = text[max(0, m.start() - 1):m.start()], text[m.end():m.end() + 8]
    if a == '-' or b.startswith('-'):
        return REV, 'palavra composta com hífen'
    if re.match(r' de mel', b):
        return REV, 'expressão "lua de mel"'
    if re.search(r'óculos de $', text[:m.start()]):
        return REV, 'expressão "óculos de sol"'
    return None, None


R(Rule('T01', 'Lua', 'Lua em maiúsculo invés do minúsculo', 'Não Substituido', r'\bluas?\b', 0,
       propose=sub_span(lambda m: m.group(0).capitalize()), ok=r'月', check=compound_check,
       regra='lua → Lua'))
R(Rule('T02', 'Sol', 'Sol em maiúsculo invés do minúsculo', 'Não Substituido', r'\b(sol|sóis)\b', 0,
       propose=sub_span(lambda m: m.group(0).capitalize()), ok=r'太陽|陽|日', ok_desc='太陽/陽/日',
       check=compound_check, regra='sol → Sol'))
R(Rule('T03', 'Olhos Místicos', 'Precisa colocar o termo com letra maiúscula.', 'Substituido',
       r'\bolhos m[ií]sticos\b', re.I, propose=canon(lambda t: 'Olhos Místicos'), ok=r'魔眼',
       regra='Resíduo: olhos místicos → Olhos Místicos'))
R(Rule('T04', 'Brasão Mágico', 'Substituir "Encravamento Mágico" por "Brasão Mágico"', 'Substituido',
       r'\bencravamentos? mágic[oa]s?\b', re.I,
       propose=sub_span(lambda m: plural(m.group(0), 'Brasão Mágico', 'Brasões Mágicos')), ok=r'魔術刻印',
       regra='Resíduo: Encravamento Mágico → Brasão Mágico'))


def encrav_check(m, text, ctx):
    jp = ctx['jp'] or ''
    if '刻印虫' in jp:
        return REV, 'resíduo de "encravamento" — JP: 刻印虫 (vermes do brasão), decidir tradução'
    if '刻印' in jp:
        return REV, 'resíduo de "encravamento" — JP: 刻印 (brasão)'
    return REV, 'resíduo de "encravamento" — JP não tem 刻印'


R(Rule('T05', 'Brasão Mágico', 'Resíduo de "encravamento" sem "Mágico" (JP 刻印)', 'Substituido',
       r'\bencravamentos?\b(?! mágic)', re.I,
       propose=sub_span(lambda m: keep_case(m.group(0), plural(m.group(0), 'brasão', 'brasões'))),
       ok=r'刻印', check=encrav_check, regra='Resíduo: encravamento → brasão (revisar)'))
R(Rule('T06', 'Brasão Mágico', 'Colocar em letra maiúscula', 'Não Substituido',
       r'\bbras(ão|ões) mágic(o|os)\b', re.I, propose=canon(title_pt), ok=r'魔術刻印|刻印', ok_desc='魔術刻印/刻印',
       regra='brasão mágico → Brasão Mágico'))
R(Rule('T07', 'Arturia', 'Substituir por "Artoria"', 'Substituido',
       r'\b(Arturia|Arthuria|Artúria|Artória|Altria)\b', 0, propose=sub_span(lambda m: 'Artoria'), ok=r'アルトリア',
       regra='Resíduo: Arturia/Altria → Artoria'))
R(Rule('T08', 'campo de separação/campo de força', 'Substituir por "Campo Limitado"', 'Substituido',
       r'\bcampos? de (separação|força)\b', re.I,
       propose=sub_span(lambda m: plural(m.group(0).split()[0], 'Campo Limitado', 'Campos Limitados')), ok=r'結界',
       regra='Resíduo: campo de separação/força → Campo Limitado'))
R(Rule('T09', 'Geas', 'Substituir "geas" por "Geas"', 'Não Substituido', r'\bgeas\b', 0,
       propose=sub_span(lambda m: 'Geas'), ok=r'ギアス', regra='geas → Geas'))
R(Rule('T10', 'Resistência Mágica', 'Substituir "resistência mágica" por "Resistência Mágica"', 'Não Substituido',
       r'\bresistências? mágicas?\b', re.I, propose=canon(title_pt), ok=r'対魔力|抗魔力', ok_desc='対魔力/抗魔力',
       regra='resistência mágica → Resistência Mágica'))
R(Rule('T10', 'Resistência Mágica', 'Variante "resistência à magia" (JP 対魔力)', 'Não Substituido',
       r'\bresistência (à|a|contra a|contra) magia\b', re.I,
       propose=sub_span(lambda m: 'Resistência Mágica'), ok=r'対魔力|抗魔力', ok_desc='対魔力/抗魔力',
       regra='resistência à magia → Resistência Mágica'))
R(Rule('T11', 'Espírito Heroico / Espíritos Heroicos', 'Colocar o termo em letra maiúscula', 'Não Substituido',
       r'\besp[ií]ritos? her[oó]icos?\b', re.I,
       propose=canon(lambda t: plural(t, 'Espírito Heroico', 'Espíritos Heroicos')), ok=r'英霊',
       regra='espírito heroico → Espírito Heroico'))
R(Rule('T12', 'Selo(s) de Comando', 'Colocar o termo em letra maiúscula', 'Não Substituido',
       r'\bselos? de comando\b', re.I, propose=canon(title_pt), ok=r'令呪', regra='selo de comando → Selo de Comando'))

FEM_MASC = {'a': 'o', 'as': 'os', 'da': 'do', 'das': 'dos', 'na': 'no', 'nas': 'nos', 'essa': 'esse', 'essas': 'esses',
            'esta': 'este', 'estas': 'estes', 'dessa': 'desse', 'desta': 'deste', 'nessa': 'nesse', 'nesta': 'neste',
            'uma': 'um', 'umas': 'uns', 'toda': 'todo', 'todas': 'todos', 'sua': 'seu', 'suas': 'seus',
            'minha': 'meu', 'minhas': 'meus', 'nossa': 'nosso', 'nossas': 'nossos', 'pela': 'pelo', 'pelas': 'pelos',
            'à': 'ao', 'às': 'aos', 'aquela': 'aquele', 'aquelas': 'aqueles'}
DET = r'(?:(\b(?:' + '|'.join(sorted(map(re.escape, FEM_MASC), key=len, reverse=True)) + r'))\s+)?'


def estampa_prop(m, text):
    det, term = m.group(1), m.group(2)
    new = plural(term.split()[0], 'Selo do Tigre', 'Selos do Tigre')
    if det:
        new = keep_case(det, FEM_MASC[det.lower()]) + ' ' + new
    return m.start(), m.end(), new


R(Rule('T13', 'Selo do Tigre', 'Encontrei algumas instâncias de "Estampa de Tigre". Padronizar como "Selo do Tigre"',
       'Não Substituido', DET + r'\b(estampas? d[eo] tigre)\b', re.I, propose=estampa_prop,
       ok=r'スタンプ', check=lambda m, t, c: (REV, 'troca de gênero (estampa→selo): conferir concordância'),
       regra='Estampa de Tigre → Selo do Tigre (+ artigo)'))


def selo_genero_prop(m, text):
    det = m.group(1)
    if not det:
        return None
    return m.start(), m.end(), keep_case(det, FEM_MASC[det.lower()]) + ' ' + m.group(2)


R(Rule('T13', 'Selo do Tigre', 'Concordância errada deixada por troca anterior (ex.: "essa Selo do Tigre")',
       'Não Substituido', DET + r'\b(Selos? do Tigre)\b', 0, propose=selo_genero_prop, ok=r'スタンプ',
       check=lambda m, t, c: (REV, 'concordância de gênero'), regra='artigo feminino + Selo do Tigre'))
R(Rule('T13', 'Selo do Tigre', 'Maiúscula no termo', 'Não Substituido', r'\bselos? do tigre\b', re.I,
       propose=canon(title_pt), ok=r'スタンプ', regra='selo do tigre → Selo do Tigre'))
R(Rule('T14', 'Dojô', '"Dojo" para "Dojô" / "dojo" para "dojô"', 'Não Substituido', r'\b[dD]ojos?\b', 0,
       propose=sub_span(lambda m: m.group(0)[:3] + 'ô' + m.group(0)[4:]), ok=r'道場', regra='dojo → dojô'))
R(Rule('T15', 'taumaturgia', 'Colocar "magia"', 'Não Substituido', r'\btaumaturgias?\b', re.I,
       propose=sub_span(lambda m: keep_case(m.group(0), plural(m.group(0), 'magia', 'magias'))),
       ok=r'魔術', block=r'魔法', regra='taumaturgia → magia'))
R(Rule('T15', 'taumaturgia', 'Derivado: taumaturgo → mago (JP 魔術師)', 'Não Substituido', r'\btaumaturg[oa]s?\b', re.I,
       propose=sub_span(lambda m: keep_case(m.group(0), 'mag' + m.group(0)[len('taumaturg'):].lower())),
       ok=r'魔術師', check=lambda m, t, c: (REV, 'derivado do termo, fora da regra literal'),
       regra='taumaturgo → mago (revisar)'))
R(Rule('T15', 'taumaturgia', 'Derivado: adjetivo taumatúrgico (fora de "campo taumatúrgico")', 'Não Substituido',
       r'(?<!campo )(?<!campos )\btaumat[uú]rgic([oa]s?)\b', re.I,
       propose=sub_span(lambda m: keep_case(m.group(0), 'mágic' + m.group(1).lower())), ok=r'魔術',
       check=lambda m, t, c: (REV, 'derivado do termo, fora da regra literal'), regra='taumatúrgico → mágico (revisar)'))
R(Rule('T16', 'Circuito Mágico / Circuitos Mágicos', 'Colocar termo em letra maiúscula', 'Não Substituido',
       r'\bcircuitos? mágicos?\b', re.I, propose=canon(title_pt), ok=r'魔術回路|回路', ok_desc='魔術回路/回路',
       regra='circuito mágico → Circuito Mágico'))
R(Rule('T17', 'Herói da Justiça', 'Encontrei algumas instâncias de "super-herói", padronizar pra "herói da justiça"',
       'Não Substituido', r'\bsuper-?her[óo]is?\b', re.I,
       propose=sub_span(lambda m: keep_case(m.group(0), plural(m.group(0), 'herói da justiça', 'heróis da justiça'))),
       ok=r'正義の味方', regra='super-herói → herói da justiça'))


def justica_check(m, text, ctx):
    if m.group(1).lower() in ('campeã', 'aliada', 'defensora', 'protetora', 'paladina'):
        return REV, 'forma feminina: decidir "heroína da justiça"?'
    return None, None


R(Rule('T17', 'Herói da Justiça', 'Outras traduções de 正義の味方 (campeão/aliado/defensor da justiça)', 'Não Substituido',
       r'\b(campeão|campeões|campeã|aliado|aliados|aliada|defensor|defensores|defensora|protetor|protetores|protetora|'
       r'paladino|paladinos|paladina) da justiça\b', re.I,
       propose=sub_span(lambda m: keep_case(m.group(0), plural(m.group(1), 'herói da justiça', 'heróis da justiça'))),
       ok=r'正義の味方', check=justica_check, regra='campeão/aliado da justiça → herói da justiça'))
R(Rule('T17', 'Herói da Justiça', 'Padronizar caixa: "herói da justiça" em minúsculo (forma majoritária no texto)',
       'Não Substituido', r'\bHer[óo]is? da [Jj]ustiça\b|\bher[óo]is? da Justiça\b', 0,
       propose=sub_span(lambda m: m.group(0).lower()), ok=r'正義の味方',
       check=lambda m, t, c: ('skip' if sentence_pos(t, m.start(), c['prev']) == 'inicio'
                              and m.group(0).split()[-1] == 'justiça' else (REV, 'padronização de caixa')),
       regra='Herói da Justiça → herói da justiça (revisar)'))


def guard_check(m, text, ctx):
    if text[max(0, m.start() - 1):m.start()] == '-':
        return REV, 'composto com hífen (ex.: contra-guardiões)'
    return None, None


R(Rule('T18', 'Guardião', 'Colocar termo em letra maiúscula', 'Não Substituido', r'\bguardi(ão|ões|ã|ãs)\b', 0,
       propose=sub_span(lambda m: m.group(0).capitalize()), ok=r'守護者|ガーディアン', ok_desc='守護者',
       check=guard_check, regra='guardião → Guardião'))
R(Rule('T19', 'Campo Limitado', 'Substituir "Campo Taumatúrgico" por "Campo Limitado"', 'Não Substituido',
       r'\bcampos? taumat[uú]rgicos?\b', re.I,
       propose=sub_span(lambda m: plural(m.group(0).split()[0], 'Campo Limitado', 'Campos Limitados')), ok=r'結界',
       regra='campo taumatúrgico → Campo Limitado'))
R(Rule('T19', 'Campo Limitado', 'Maiúscula no termo', 'Não Substituido', r'\bcampos? limitados?\b', re.I,
       propose=canon(title_pt), ok=r'結界', regra='campo limitado → Campo Limitado'))


def magia_check(m, text, ctx):
    if re.match(r' Verdadeiras?\b', text[m.end():]):
        return 'skip'
    pos = sentence_pos(text, m.start(), ctx['prev'])
    if pos == 'inicio':
        ctx['stats']['magia_inicio_frase'] += 1
        return 'skip'
    if pos == 'ambiguo':
        return REV, 'depois de reticências: conferir se é início de frase'
    return None, None


R(Rule('T20', '"magia"', 'Casos em que "magia" (magecraft) está como "Magia". Padronizar como "magia", a não ser no '
       'começo de frase', 'Não Substituido', r'\bMagias?\b', 0,
       propose=sub_span(lambda m: m.group(0).lower()), ok=r'魔術', block=r'魔法', ok_desc='魔術 (majutsu)',
       block_desc='魔法 (mahou = Magia Verdadeira)', check=magia_check, regra='Magia → magia (meio de frase)'))

# --- Extra fora da planilha (Discord): energia mágica
EXTRA = [Rule('X01', 'Energia Mágica (Discord)', 'DAISO: "Energia Mágica" em maiúscula sem necessidade. '
              'ATENÇÃO: o Glossário TYPE-MOON lista 魔力 = "Energia Mágica"', '—',
              r'\b[Ee]nergias? [Mm]ágicas?\b', 0, propose=None, ok=r'魔力', regra='Energia Mágica → energia mágica')]


def energia_prop(m, text, prev):
    t = m.group(0)
    start = sentence_pos(text, m.start(), prev) == 'inicio'
    new = t.lower()
    if start:
        new = new[0].upper() + new[1:]
    return None if new == t else (m.start(), m.end(), new)


# --- Cobertura JP: termo JP presente mas PT sem a forma padrão
COBERTURA = [
    ('正義の味方', r'正義の味方', r'her[óo]i(s)? da justiça|heroína da justiça', 'herói da justiça'),
    ('令呪', r'令呪', r'selos? de comando', 'Selo de Comando'),
    ('英霊', r'英霊', r'esp[ií]ritos? her[oó]ic', 'Espírito Heroico'),
    ('魔術回路', r'魔術回路', r'circuitos? mágic', 'Circuito Mágico'),
    ('魔術刻印', r'魔術刻印', r'bras(ão|ões) mágic', 'Brasão Mágico'),
    ('結界', r'(?<!固有)結界', r'campos? limitad|campos? taumat|campos? de (separação|força)', 'Campo Limitado'),
    ('対魔力/抗魔力', r'対魔力|抗魔力', r'resistências? mágica|resistência (à|a) magia', 'Resistência Mágica'),
    ('魔眼', r'魔眼', r'olhos m[ií]sticos', 'Olhos Místicos'),
    ('道場', r'道場', r'\bdoj[oô]', 'dojô'),
    ('ギアス', r'ギアス', r'\bgeas\b', 'Geas'),
    ('守護者', r'守護者', r'guardi', 'Guardião'),
    ('アルトリア', r'アルトリア', r'artoria', 'Artoria'),
    ('魔法', r'魔法', r'Magias? Verdadeiras?|\bMagias?\b', 'Magia Verdadeira'),
]


# ---------------------------------------------------------------- varredura

def main():
    prog = load_progresso()
    stats = Counter()
    rows, extra_rows, cob_rows, ver_rows = [], [], [], []

    files = sorted(f for f in os.listdir(PT_DIR) if f.endswith('.epk_dec'))
    for f in files:
        pt, anom = parse_dat(os.path.join(PT_DIR, f))
        jp, anom_jp = parse_dat(os.path.join(JP_DIR, f))
        meta = prog.get(f, {'cena': '', 'revisores': '', 'estado': '', 'rota': 'Dados/Sistema'})
        stats['arquivos'] += 1
        if not os.path.exists(os.path.join(JP_DIR, f)):
            stats['sem_jp'] += 1
        prev = None
        audit_keys = defaultdict(set)
        for key, (ln, text) in sorted(pt.items(), key=lambda kv: kv[1][0]):
            stats['linhas'] += 1
            jtext = jp.get(key, (None, None))[1]
            jln = jp.get(key, (None, None))[0]
            if jtext is not None:
                stats['linhas_pareadas'] += 1
            base = dict(arquivo=f, linha=ln, label=key[0] + ('' if key[1] in ('text', 'name') else f' [{key[1]}]'),
                        linha_jp=jln, pt=text, jp=jtext, rota=meta['rota'], cena=meta['cena'],
                        estado_script=meta['estado'], hash=hashlib.sha1(text.encode('utf-8')).hexdigest()[:12])
            ctx = {'jp': jtext, 'prev': prev, 'stats': stats}

            for rule in RULES:
                for m in rule.rx.finditer(text):
                    p = rule.propose(m, text)
                    if p is None:
                        continue
                    forced, motivo_f = None, None
                    if rule.check:
                        r = rule.check(m, text, ctx)
                        if r == 'skip':
                            continue
                        forced, motivo_f = r
                    s, e, new = p
                    cls, motivo = classify(jtext, rule.ok, rule.block, rule.ok_desc, rule.block_desc)
                    if in_tag(text, m.start()):
                        forced, motivo_f = REV, 'ocorrência dentro de uma tag [...]'
                    if forced and not (cls == BLOQ):
                        motivo = f'{motivo_f}; {motivo}'
                        cls = forced
                    elif forced:
                        motivo = f'{motivo_f}; {motivo}'
                    rows.append(dict(base, rid=rule.rid, termo=rule.termo, obs=rule.obs, estado_planilha=rule.estado,
                                     regra=rule.regra, classe=cls, motivo=motivo, s=s, e=e,
                                     achado=text[s:e], novo=new, pt_novo=text[:s] + new + text[e:]))
                    audit_keys[rule.termo].add(key)

            for rule in EXTRA:
                for m in rule.rx.finditer(text):
                    p = energia_prop(m, text, prev)
                    if p is None:
                        continue
                    s, e, new = p
                    cls, motivo = classify(jtext, rule.ok, None, rule.ok, '')
                    extra_rows.append(dict(base, rid=rule.rid, termo=rule.termo, obs=rule.obs, estado_planilha='—',
                                           regra=rule.regra, classe=cls, motivo=motivo, s=s, e=e, achado=text[s:e],
                                           novo=new, pt_novo=text[:s] + new + text[e:]))

            if jtext is not None:
                for nome, jrx, ptrx, forma in COBERTURA:
                    if re.search(jrx, jtext) and not re.search(ptrx, text, re.I if nome != '魔法' else 0):
                        cob_rows.append(dict(base, termo=nome, forma=forma, classe=REV,
                                             motivo=f'JP tem {nome}, PT não usa "{forma}"'))
                if '燕返し' in jtext or re.search(r'tsubame|andorinha', text, re.I):
                    ver_rows.append(dict(base, termo='Tsubame Gaeshi', classe=DEC,
                                         motivo='Vai manter o NP do Assassin dessa forma ou vai traduzir?', novo=''))
                if 'トレース' in jtext or re.search(r'\btrace\b', text, re.I):
                    ver_rows.append(dict(base, termo='Ruby "Trace" do Shirou', classe=REV,
                                         motivo='Verificar ruby texts dos "Trace" do Shirou', novo=''))
            if '♪' in text or (jtext and '♪' in jtext):
                novo = text.replace('♪[humming]', '[humming]').replace('♪', '[humming]') if '♪' in text else ''
                ver_rows.append(dict(base, termo='Caractere ♪', classe=REV,
                                     motivo='Caractere bugado; pode ser substituído por [humming]'
                                            + ('' if '♪' in text else ' (♪ só no JP)'), novo=novo))
            prev = text

    write_xlsx(rows, extra_rows, cob_rows, ver_rows, stats)


# ---------------------------------------------------------------- planilha

FILL = {AUTO: PatternFill('solid', fgColor='D9EAD3'), REV: PatternFill('solid', fgColor='FFF2CC'),
        BLOQ: PatternFill('solid', fgColor='F4CCCC'), DEC: PatternFill('solid', fgColor='CFE2F3')}
HEAD_FILL = PatternFill('solid', fgColor='434343')
HEAD_FONT = Font(bold=True, color='FFFFFF')
RED = InlineFont(b=True, color='CC0000')
GREEN = InlineFont(b=True, color='38761D')


def clean(s):
    if s is None:
        return ''
    return ILLEGAL_CHARACTERS_RE.sub('', s)[:32000]


def rich(text, s, e, font):
    text = clean(text)
    parts = [p for p in (text[:s], TextBlock(font, text[s:e]) if e > s else None, text[e:]) if p]
    return CellRichText(*parts) if parts else ''


def vscode(pasta, arq, ln):
    return 'vscode://file/' + urllib.parse.quote(os.path.join(pasta, arq).replace('\\', '/'), safe='/:') + f':{ln}'


def links(r):
    pt = vscode(PT_DIR, r['arquivo'], r['linha'])
    jp = vscode(JP_DIR, r['arquivo'], r['linha_jp']) if r.get('linha_jp') else None
    return pt, jp


def setup(ws, headers, widths):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 30


def put_link(cell, url, label):
    if url:
        cell.value, cell.hyperlink = label, url
        cell.font = Font(color='1155CC', underline='single')


def write_occ_sheet(ws, rows, with_approve=True):
    headers = ['ID', 'Termo (planilha)', 'Regra', 'Classificação', 'Motivo', 'Aprovar (S/N)', 'Rota', 'Cena',
               'Estado do script', 'Arquivo', 'Linha PT', 'Linha JP', 'Label', 'Abrir PT', 'Abrir JP',
               'Encontrado', 'Proposta', 'PT atual', 'PT proposto', 'JP', 'Obs (planilha)', 'Hash PT', 'Posição']
    widths = [11, 22, 30, 13, 40, 9, 9, 30, 12, 34, 8, 8, 30, 9, 9, 20, 20, 70, 70, 60, 40, 13, 8]
    setup(ws, headers, widths)
    wrap = Alignment(wrap_text=True, vertical='top')
    top = Alignment(vertical='top')
    for i, r in enumerate(rows, start=1):
        lpt, ljp = links(r)
        ws.append([f"{r['rid']}-{i:05d}", r['termo'], r['regra'], r['classe'], r['motivo'],
                   'S' if r['classe'] == AUTO else '', r['rota'], clean(r['cena']), r['estado_script'], r['arquivo'],
                   r['linha'], r['linha_jp'], r['label'], None, None, clean(r['achado']), clean(r['novo']), None, None,
                   clean(r['jp'] if r['jp'] is not None else '(sem par no JP)'), r['obs'], r['hash'], r['s']])
        n = ws.max_row
        ws.cell(n, 18).value = rich(r['pt'], r['s'], r['e'], RED)
        ws.cell(n, 19).value = rich(r['pt_novo'], r['s'], r['s'] + len(r['novo']), GREEN)
        put_link(ws.cell(n, 14), lpt, 'PT ↗')
        put_link(ws.cell(n, 15), ljp, 'JP ↗')
        ws.cell(n, 4).fill = FILL[r['classe']]
        for c in range(1, len(headers) + 1):
            ws.cell(n, c).alignment = wrap if c in (5, 8, 18, 19, 20, 21) else top
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{ws.max_row}'
    if with_approve and rows:
        dv = DataValidation(type='list', formula1='"S,N"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'F2:F{ws.max_row}')


def write_simple_sheet(ws, rows, extra_col=None):
    headers = ['Termo', 'Classificação', 'Motivo', 'Rota', 'Cena', 'Estado do script', 'Arquivo', 'Linha PT',
               'Linha JP', 'Label', 'Abrir PT', 'Abrir JP', 'PT atual', 'JP']
    widths = [22, 13, 40, 9, 30, 12, 34, 8, 8, 30, 9, 9, 70, 60]
    if extra_col:
        headers.insert(13, extra_col)
        widths.insert(13, 60)
    setup(ws, headers, widths)
    wrap = Alignment(wrap_text=True, vertical='top')
    for r in rows:
        lpt, ljp = links(r)
        vals = [r['termo'], r['classe'], r['motivo'], r['rota'], clean(r['cena']), r['estado_script'], r['arquivo'],
                r['linha'], r['linha_jp'], r['label'], None, None, clean(r['pt']),
                clean(r['jp'] or '(sem par no JP)')]
        if extra_col:
            vals.insert(13, clean(r.get('novo') or r.get('forma', '')))
        ws.append(vals)
        n = ws.max_row
        put_link(ws.cell(n, 11), lpt, 'PT ↗')
        put_link(ws.cell(n, 12), ljp, 'JP ↗')
        ws.cell(n, 2).fill = FILL[r['classe']]
        for c in range(1, len(headers) + 1):
            ws.cell(n, c).alignment = wrap
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{ws.max_row}'


def write_xlsx(rows, extra_rows, cob_rows, ver_rows, stats):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumo'
    bold = Font(bold=True)
    info = [
        ('Auditoria de termos — FSN Remastered PT-BR', Font(bold=True, size=14)),
        (f'Base local: {PT_DIR}  ×  {JP_DIR}', None),
        (f'Gerada em {datetime.now():%d/%m/%Y %H:%M}. Links "PT ↗"/"JP ↗" abrem o arquivo local na linha exata '
         '(VS Code). Se o script for editado depois, gere de novo.', None),
        (f"{stats['arquivos']} arquivos, {stats['linhas']} linhas de texto, {stats['linhas_pareadas']} pareadas com o JP "
         'pelo mesmo label.', None),
        (f"\"Magia\" em início de frase ignorado (correto): {stats['magia_inicio_frase']}.", None),
        ('', None),
        ('Como usar', bold),
        ('• Aba "Auditoria": uma linha por ocorrência. Nenhum script foi alterado.', None),
        ('• AUTO = o JP da mesma linha confirma o termo → "Aprovar" já vem S. Só marque N se estiver errado.', None),
        ('• REVISAR = ambíguo ou o JP não confirma → decida S ou N.', None),
        ('• BLOQUEADO = o JP contradiz a troca (ex.: 魔法 = Magia Verdadeira) → não trocar, a não ser que marque S.', None),
        ('• "PT atual" destaca em vermelho o trecho encontrado; "PT proposto" destaca em verde a troca.', None),
        ('• "Hash PT" e "Posição" servem pra aplicar depois com segurança: se a linha mudou desde a auditoria, '
         'a troca é recusada.', None),
        ('• Aba "Cobertura JP": linhas onde o JP tem o termo mas o PT não usa a forma padrão (traduções inconsistentes '
         'ou pronome/omissão legítima). Só leitura/decisão.', None),
        ('• Aba "Verificações": Tsubame Gaeshi, caractere ♪ e ruby do "Trace".', None),
        ('• Aba "Extra - Energia Mágica": veio do Discord, NÃO está na planilha de termos, e o glossário diz '
         '"Energia Mágica" com maiúscula. Decidir antes.', None),
        ('', None),
    ]
    for text, font in info:
        ws.append([text])
        if font:
            ws.cell(ws.max_row, 1).font = font
    ws.append([])
    hdr = ['Termo (planilha)', 'Regra', 'Estado na planilha', AUTO, REV, BLOQ, 'Total']
    ws.append(hdr)
    hr = ws.max_row
    for c in range(1, len(hdr) + 1):
        ws.cell(hr, c).fill, ws.cell(hr, c).font = HEAD_FILL, HEAD_FONT
    summary = defaultdict(Counter)
    order = []
    for r in rows + extra_rows:
        k = (r['termo'], r['regra'], r['estado_planilha'])
        if k not in summary:
            order.append(k)
        summary[k][r['classe']] += 1
    seen = set()
    for rule in RULES + EXTRA:  # mostra também regras sem nenhuma ocorrência
        k = (rule.termo, rule.regra, rule.estado)
        if k not in summary and k not in seen:
            order.append(k)
        seen.add(k)
    for k in order:
        c = summary[k]
        ws.append([k[0], k[1], k[2], c[AUTO], c[REV], c[BLOQ], sum(c.values())])
        for col, cls in ((4, AUTO), (5, REV), (6, BLOQ)):
            if c[cls]:
                ws.cell(ws.max_row, col).fill = FILL[cls]
    ws.append([])
    ws.append(['Cobertura JP (linhas)', '', '', '', len(cob_rows)])
    for k, v in Counter(r['termo'] for r in cob_rows).most_common():
        ws.append(['', k, '', '', v])
    ws.append(['Verificações (linhas)', '', '', '', len(ver_rows)])
    for k, v in Counter(r['termo'] for r in ver_rows).most_common():
        ws.append(['', k, '', '', v])
    for col, w in zip('ABCDEFG', (38, 48, 18, 10, 10, 12, 8)):
        ws.column_dimensions[col].width = w

    write_occ_sheet(wb.create_sheet('Auditoria'), rows)
    write_simple_sheet(wb.create_sheet('Cobertura JP'), cob_rows, extra_col='Forma padrão esperada')
    write_simple_sheet(wb.create_sheet('Verificações'), ver_rows, extra_col='Proposta')
    write_occ_sheet(wb.create_sheet('Extra - Energia Mágica'), extra_rows)
    wb.save(OUT)

    print('OK ->', OUT)
    print(dict(stats))
    print('Auditoria:', len(rows), Counter(r['classe'] for r in rows))
    for k in order:
        print(f'  {sum(summary[k].values()):5}  {dict(summary[k])}  {k[0]} :: {k[1]}')
    print('Cobertura JP:', len(cob_rows), dict(Counter(r['termo'] for r in cob_rows)))
    print('Verificações:', len(ver_rows), dict(Counter(r['termo'] for r in ver_rows)))
    print('Extra:', len(extra_rows), dict(Counter(r['classe'] for r in extra_rows)))


if __name__ == '__main__':
    main()
